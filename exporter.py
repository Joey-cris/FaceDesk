"""Attendance export to Excel (.xlsx) or CSV."""
from __future__ import annotations

import csv
from datetime import datetime

from db import TS

COLUMNS = ["Date", "Name", "Student ID", "Section", "Role", "Time In", "Time Out", "Hours", "Status", "Source", "Note"]
SUMMARY_COLUMNS = ["Name", "Student ID", "Section", "School days", "Present", "Late", "Absent", "Attendance %"]


def _t(value) -> str:
    return datetime.strptime(value, TS).strftime("%I:%M %p").lstrip("0") if value else ""


def to_rows(records: list[dict]) -> list[list]:
    rows = []
    for r in records:
        hours = ""
        if r["time_out"]:
            delta = datetime.strptime(r["time_out"], TS) - datetime.strptime(r["time_in"], TS)
            hours = round(delta.total_seconds() / 3600, 2)
        rows.append([r["day"], r["name"], r.get("student_id", ""), r.get("section", ""), r["role"],
                     _t(r["time_in"]), _t(r["time_out"]), hours, "Late" if r["late"] else "On time",
                     "Manual" if r.get("source") == "manual" else "Face", r.get("note", "")])
    return rows


def summary_rows(summary: list[dict]) -> list[list]:
    return [[r["name"], r["student_id"], r["section"], r["school_days"], r["present"], r["late"],
             r["absent"], r["rate"]] for r in summary]


def _write_csv(columns, rows, path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: Excel opens it cleanly
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(rows)


def _write_xlsx(columns, rows, path, title, late_col=None) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel export needs 'openpyxl'. Run setup.bat again, or export CSV instead.") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A4A7A")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
        if late_col is not None and row[late_col] == "Late":
            ws.cell(ws.max_row, late_col + 1).font = Font(color="C0392B", bold=True)
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 3, 10), 40)
    ws.freeze_panes = "A2"
    wb.save(str(path))


def export_csv(records: list[dict], path) -> None:
    _write_csv(COLUMNS, to_rows(records), path)


def export_xlsx(records: list[dict], path) -> None:
    _write_xlsx(COLUMNS, to_rows(records), path, "Attendance", late_col=COLUMNS.index("Status"))


def export_summary_csv(summary: list[dict], path) -> None:
    _write_csv(SUMMARY_COLUMNS, summary_rows(summary), path)


def export_summary_xlsx(summary: list[dict], path) -> None:
    _write_xlsx(SUMMARY_COLUMNS, summary_rows(summary), path, "Summary")


def export_audit_csv(entries: list[dict], path) -> None:
    _write_csv(["Time", "Who", "Action", "Detail"], [[e["ts"], e["actor"], e["action"], e["detail"]] for e in entries], path)


TRIAL_COLUMNS = ["Session", "Time", "Kind", "Truth", "Lighting", "Distance", "Accessory", "Note", "Trial", "Outcome",
                 "Accepted as (id)", "Correct", "Closest dist", "Closest id", "Runner-up", "Seconds",
                 "Liveness blocked", "Threshold", "Margin"]


def trial_rows(trials: list[dict]) -> list[list]:
    r = lambda v: "" if v is None else round(v, 2)
    return [[t["session"], t["ts"], t["truth_kind"], t["truth_label"], t["lighting"], t["distance"], t["accessory"],
             t["note"], t["trial_no"], t["outcome"], t["accepted_person_id"] if t["accepted_person_id"] is not None else "",
             "Yes" if t["correct"] else "No", r(t["min_dist"]), t["min_label"] if t["min_label"] is not None else "",
             r(t["runner_up"]), r(t["seconds"]), "Yes" if t["liveness_blocked"] else "", r(t["threshold"]), r(t["margin"])]
            for t in trials]


def export_trials_csv(trials: list[dict], path) -> None:
    _write_csv(TRIAL_COLUMNS, trial_rows(trials), path)


def export_accuracy_xlsx(trials: list[dict], path, thresholds=range(40, 95, 5)) -> None:
    """Workbook for Chapter 4: Summary (overall + by condition), Trials (raw data), Threshold analysis."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Excel export needs 'openpyxl'. Run setup.bat again, or export CSV instead.") from exc
    import metrics
    wb = Workbook()
    head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="0A4A7A"))
                       or setattr(c, "alignment", Alignment(horizontal="center")) for c in ws[1]]
    ws = wb.active
    ws.title = "Summary"
    cols = ["Group", "Value", "Attempts", "Genuine", "Impostor", "TAR %", "FRR %", "Misid %", "FAR %", "TRR %",
            "Accuracy %", "Mean sec (correct accepts)"]
    ws.append(cols); head(ws)
    def line(group, value, m):
        v = lambda x: "" if x is None else x
        ws.append([group, value, m["attempts"], m["genuine"], m["impostor"], v(m["TAR"]), v(m["FRR"]), v(m["MISID"]),
                   v(m["FAR"]), v(m["TRR"]), v(m["accuracy"]), v(m["mean_seconds"])])
    line("Overall", "all trials", metrics.summarize(trials))
    for key in ("lighting", "distance", "accessory"):
        for row in metrics.by_condition(trials, key):
            line(key.title(), row["value"], row)
    ws.append([])
    ws.append(["Definitions: TAR = correct accepts / genuine attempts. FRR = genuine not accepted / genuine. "
               "Misid = genuine accepted as another person / genuine. FAR = impostors accepted / impostor attempts. "
               "Accuracy = (correct accepts + correct rejects) / all attempts."])
    t = wb.create_sheet("Trials")
    t.append(TRIAL_COLUMNS); head(t)
    for row in trial_rows(trials):
        t.append(row)
    th = wb.create_sheet("Threshold analysis")
    th.append(["Threshold", "TAR % (approx)", "FRR % (approx)", "FAR % (approx)", "Genuine", "Impostor"]); head(th)
    for r_ in metrics.threshold_table(trials, thresholds):
        th.append([r_["threshold"], r_["TAR"] if r_["TAR"] is not None else "", r_["FRR"] if r_["FRR"] is not None else "",
                   r_["FAR"] if r_["FAR"] is not None else "", r_["genuine"], r_["impostor"]])
    th.append([]); th.append(["Approximate: uses only each trial's closest-match distance (ignores the margin, "
                              "consecutive-frame and liveness rules). Confirm the chosen threshold with a new test."])
    for sheet in (ws, t, th):
        for col in sheet.columns:
            w = max((len(str(c.value)) if c.value is not None else 0) for c in col[:60])
            sheet.column_dimensions[col[0].column_letter].width = min(max(w + 2, 9), 34)
        sheet.freeze_panes = "A2"
    wb.save(str(path))
